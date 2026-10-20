"""Extrae SUMMARY y las pestañas MK HQ, MKDIST y MKBSTORES del Excel Marketing Q4 a JSON (lo usa generar.py)."""
import json
import sys

import openpyxl

PESTANAS = {"hq": "MK HQ", "dist": " MKDIST", "bstores": "MKBSTORES"}
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)  # read_only: la caché de tablas dinámicas es enorme
num = lambda v: round(float(v), 2) if isinstance(v, (int, float)) else 0.0


def leer(nombre: str) -> dict:
    filas = [list(r) + [None] * 10 for r in wb[nombre].iter_rows(values_only=True)]
    # La tabla dinámica empieza en la fila "Etiquetas de fila"; la columna de esa celda es la del proveedor
    ini, col = next((i, r.index("Etiquetas de fila")) for i, r in enumerate(filas) if "Etiquetas de fila" in r)
    filtros = {r[col]: r[col + 1] for r in filas[:ini] if r[col] in ("MES", "BUN", "TRESD")}
    proveedores = []
    for r in filas[ini + 1:]:
        prov = r[col]
        if not prov or prov in ("Total general", "HQ") or (prov == "(en blanco)" and abs(num(r[col + 1])) < 1):
            continue
        p = {"proveedor": str(prov).strip(), "importe": num(r[col + 1]), "seguimiento": num(r[col + 2]),
             "comentario": str(r[col + 4]).strip() if col + 4 < len(r) and isinstance(r[col + 4], str) else ""}
        tipo = r[col - 1] if col > 0 else None
        p["tipo"] = tipo if tipo in ("BRAND", "DIGITAL") else ("Ajustes" if "AJUSTES" in p["proveedor"] else "Sin tipo")
        proveedores.append(p)
    return {"filtros": {k: str(v) for k, v in filtros.items()}, "proveedores": proveedores}


def leer_resumen() -> dict:
    """Tabla de SUMMARY (miles de euros): una fila por unidad y la fila TOTAL; columnas desde la cabecera con BUDGET."""
    filas = [list(r) + [None] * 10 for r in wb["SUMMARY"].iter_rows(values_only=True)]
    ini = next(i for i, r in enumerate(filas) if "BUDGET" in r)
    cab = filas[ini]
    cols = [(j, str(v).strip()) for j, v in enumerate(cab) if isinstance(v, str) and v.strip()]
    unidades = []
    for r in filas[ini + 1:]:
        etiqueta = next((v for v in r[:cols[0][0]] if isinstance(v, str) and v.strip()), None)
        if not etiqueta:
            continue
        unidades.append({"unidad": etiqueta.strip(), "valores": {nombre: num(r[j]) for j, nombre in cols}})
        if etiqueta.strip().upper() == "TOTAL":
            break
    return {"columnas": [nombre for _, nombre in cols], "filas": unidades}


salida = {k: leer(v) for k, v in PESTANAS.items()}
salida["resumen"] = leer_resumen()
json.dump(salida, sys.stdout, ensure_ascii=False, separators=(",", ":"))
