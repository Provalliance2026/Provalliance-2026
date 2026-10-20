"""Extrae Seguimiento, TARJETA y TOTAL2026 (TravelPerk) del Excel de viajes a JSON (lo usa generar.py)."""
import collections, json, sys
import openpyxl

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
num = lambda v: round(float(v) * 1000, 2) if isinstance(v, (int, float)) else None  # k€ -> €
txt = lambda v: str(v).strip() if v not in (None, "") else ""

# ── Seguimiento (k€ en el Excel; se pasa a €) ──
seg = [list(r) + [None] * 40 for r in wb["Seguimiento"].iter_rows(values_only=True)]
resumen = []
for r in seg[5:16]:  # filas 6-16
    if r[4]:
        resumen.append({"owner": txt(r[3]), "bu": txt(r[4]), "total": not r[3] or r[4] in ("TOTAL",),
                        "presupuesto": num(r[6]), "meses": [num(r[c]) for c in range(9, 21)],
                        "real": num(r[22]), "pendiente": num(r[25])})
grupos = {22: "Hairdressing · Jaime Marrero", 33: "Training · Patrick Philipeu",
          43: "Salon Space · Manuel Estevez", 59: "HQ", 66: "Education"}
detalle, grupo = [], None
for i in range(22, 69):
    r = seg[i - 1]
    if i in grupos:
        grupo = grupos[i]; continue
    if r[2] is not None or not r[4]:  # filas de reparto (0,5 / 0,25…) y vacías
        continue
    detalle.append({"grupo": grupo, "bu": txt(r[3]), "persona": txt(r[4]), "presupuesto": num(r[6]),
                    "meses": [num(r[c]) for c in range(9, 21)],
                    "real": num(r[22]) if r[22] is not None else round(sum(num(r[c]) or 0 for c in range(9, 21)), 2),
                    "pendiente": num(r[25]), "tendencia": num(r[26]), "desv": num(r[27]),
                    "desvPct": r[28] if isinstance(r[28], (int, float)) else None, "tarjeta": num(r[31])})
tot = seg[69]
total_det = {"presupuesto": num(tot[6]), "meses": [num(tot[c]) for c in range(9, 21)], "real": num(tot[22]),
             "pendiente": num(tot[25]), "tarjeta": num(tot[31])}
fuera = [{"bu": txt(seg[i-1][3]), "persona": txt(seg[i-1][4]), "meses": [num(seg[i-1][c]) for c in range(9, 21)],
          "real": num(seg[i-1][22])} for i in (77, 78)]

# ── TARJETA (€) ──
tar = [list(r) + [None] * 30 for r in wb["TARJETA"].iter_rows(values_only=True)]
meses_tar = MESES[:9]
tarjeta = []
for r in tar[1:38]:
    if r[9]:
        tarjeta.append({"titular": txt(r[9]).title(), "bun": txt(r[10]), "cruce": txt(r[11]),
                        "meses": [round(float(r[c]), 2) if isinstance(r[c], (int, float)) else 0 for c in range(13, 22)]})

# ── TOTAL2026 TravelPerk (€), agregado por mes/BU/sub-BU/cruce/empleado/naturaleza ──
agg = collections.defaultdict(lambda: [0.0, 0])
for r in list(wb["TOTAL2026"].iter_rows(values_only=True))[1:]:
    if r[11] is None: continue
    sub = r[9] if r[9] not in (None, 0, "0") else "Sin sub-BU"
    k = (MESES.index(r[6]), txt(r[8]) or "Sin BU", txt(sub), txt(r[10]), " ".join(txt(r[7]).split()).title(),
         txt(r[14]) or "Sin clasificar")
    agg[k][0] += float(r[11]); agg[k][1] += 1
travelperk = [[*k, round(v[0], 2), v[1]] for k, v in sorted(agg.items())]

json.dump({"meses": MESES, "mesesTarjeta": meses_tar,
           "seguimiento": {"resumen": resumen, "detalle": detalle, "total": total_det, "fueraTotal": fuera},
           "tarjeta": tarjeta, "travelperk": travelperk}, sys.stdout, ensure_ascii=False, separators=(",", ":"))
